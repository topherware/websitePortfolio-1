from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('C:/Web Gacor')
OUT = ROOT / 'artifacts' / 'Panduan Penjualan Portfolio.docx'
doc = Document()
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.7)
s.left_margin = s.right_margin = Inches(.8)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name, size in [('Title', 27), ('Heading 1', 19), ('Heading 2', 13)]:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string('000000' if name == 'Title' else '17324D')
    st.paragraph_format.space_before = Pt(12)
    st.paragraph_format.space_after = Pt(8)
doc.core_properties.title = 'Panduan Penjualan Portfolio'
doc.core_properties.subject = 'Prosedur setup Vercel dan Supabase untuk setiap pembeli'
doc.core_properties.author = 'Portfolio Admin'

def p(text, style=None): return doc.add_paragraph(text, style)
def h(text): doc.add_heading(text, 2)
def step(n, text): p(f'{n}. {text}')
def code(text):
    para = p(text)
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after = Pt(10)
    for run in para.runs:
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)
def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for i, v in enumerate(headers): t.rows[0].cells[i].text = v
    for row in rows:
        for c, val in zip(t.add_row().cells, row): c.text = val
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri == 0: trpr.append(OxmlElement('w:tblHeader'))
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            tcpr = cell._tc.get_or_add_tcPr()
            sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), '17324D' if ri == 0 else ('F0F4F7' if ri % 2 else 'FFFFFF')); tcpr.append(sh)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:' + edge); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D9D9D9'); borders.append(el)
            tcpr.append(borders)
            mar = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:' + edge); el.set(qn('w:w'), '90'); el.set(qn('w:type'), 'dxa'); mar.append(el)
            tcpr.append(mar)
            va = OxmlElement('w:vAlign'); va.set(qn('w:val'), 'center'); tcpr.append(va)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.0
                for r in para.runs:
                    r.font.size = Pt(10.5)
                    if ri == 0: r.bold = True; r.font.color.rgb = RGBColor(255,255,255)
    p('')

p('Panduan Penjualan Portfolio', 'Title')
p('Setup setiap pembeli dari pembayaran sampai serah terima', 'Subtitle')
p('Versi 1  |  2 Oktober 2026')
p('Gunakan panduan ini setiap kali ada pembeli baru. Kamu menyiapkan website, database, dan undangan; pembeli membuat password sendiri melalui email. Jangan meminta password Gmail atau password admin pembeli.')
h('Model layanan yang digunakan')
p('Satu pembeli memakai satu project Vercel dan satu project Supabase. Struktur aplikasi saat ini hanya menyimpan satu portfolio dengan ID site. Semua anggota admin_users pada project yang sama dapat mengelola konten tersebut.')
p('Project tetap berada di akunmu. Sepakati biaya hosting, domain, pemeliharaan, backup, dan cara pemindahan kepemilikan sejak awal. Periksa ketentuan penggunaan komersial serta kuota paket Vercel dan Supabase sebelum menjual layanan; jangan menjanjikan hosting gratis tanpa batas.')
h('1 Catat pesanan dan siapkan salinan')
step(1, 'Minta nama pembeli, email login, nama portfolio, domain yang diinginkan, dan konten awal.')
step(2, 'Buat folder atau repository privat khusus pembeli dari template terbaru, misalnya portfolio-budi. Jangan menyalin data atau konfigurasi pembeli sebelumnya.')
step(3, 'Pastikan tersedia akses Vercel, Supabase, repository, layanan email pengirim, dan pengelolaan DNS jika memakai domain sendiri.')
table(['Tahap', 'Hasil yang harus tersedia'], [
('Database', 'Schema, bucket portfolio, dan satu record konten site'),
('Deployment', 'Website publik dan halaman /admin/ terbuka'),
('Autentikasi', 'Signup publik mati, email aktif, redirect dan SMTP benar'),
('Aktivasi', 'Undangan terkirim dan UUID pembeli diberi izin admin'),
('Serah terima', 'Login, simpan, upload, dan reset password sudah diuji')], [1.45,5.45])
p('Nama menu dashboard dapat berubah. Ikuti fungsi pengaturannya jika label yang tampil berbeda.')

page('2 Siapkan database dan konten awal')
step(1, 'Di Supabase, pilih New project. Beri nama sesuai pembeli, pilih region, lalu simpan password database secara aman. Password ini bukan password login admin pembeli.')
step(2, 'Buka SQL Editor dan jalankan isi supabase-schema.sql pada project baru. Script membuat tabel, view publik, dan aturan akses RLS.')
step(3, 'Buka Storage dan buat bucket bernama portfolio. Aktifkan Public agar aset website dapat diakses melalui URL. Jangan unggah dokumen rahasia ke bucket ini.')
h('Isi record site sebelum pembeli mengedit')
p('Schema belum memasukkan data awal. Gunakan data contoh dari cms.js, lalu ganti dengan konten pembeli melalui admin. Jangan memakai JSON kosong karena editor membutuhkan struktur konten lengkap.')
step(4, 'Di folder utama salinan proyek, buat file export-seed.mjs berisi kode berikut. Node.js harus tersedia di komputer.')
code('''import { writeFileSync } from "node:fs";
globalThis.window = {};
const { demoContent } = await import("./cms.js");
writeFileSync("portfolio-seed.json",
  JSON.stringify(demoContent, null, 2));''')
step(5, 'Buka terminal pada folder tersebut, lalu jalankan:')
code('node export-seed.mjs')
step(6, 'Buka Supabase Table Editor, pilih portfolio_content, lalu Insert row. Isi id dengan site. Salin seluruh isi portfolio-seed.json ke kolom JSON content, biarkan updated_at memakai default, lalu simpan.')
step(7, 'Pastikan ada tepat satu record dengan id site dan content berisi objek JSON lengkap. Jika record site sudah ada, jangan menimpanya tanpa memeriksa data yang tersimpan.')
h('Periksa sebelum lanjut')
p('Tabel admin_users masih boleh kosong. Tabel portfolio_content harus sudah berisi site, dan bucket portfolio harus tersedia. RLS tetap aktif pada tabel; jangan mematikannya untuk mengatasi masalah login.')
p('File export-seed.mjs dan portfolio-seed.json hanya alat bantu lokal. Hasil build aplikasi tidak memerlukan kedua file tersebut.')

page('3 Hubungkan website dan atur autentikasi')
h('Konfigurasi dan deployment')
step(1, 'Ambil Project URL dan anon key project Supabase pembeli. Isi config.js mengikuti config.example.js:')
code('''window.__PORTFOLIO_CONFIG__ = {
  supabaseUrl: "https://PROJECT-PEMBELI.supabase.co",
  supabaseAnonKey: "ANON_KEY_PROJECT_PEMBELI",
  storageBucket: "portfolio",
};''')
p('File ini dibaca publik. Jangan memasukkan service_role, secret key, password database, atau kredensial SMTP di dalamnya.')
step(2, 'Push perubahan ke repository pembeli, lalu import sebagai project baru di Vercel. Gunakan pengaturan berikut:')
table(['Pengaturan Vercel', 'Nilai'], [('Framework Preset','Other'),('Build Command','npm run build'),('Output Directory','dist')], [2.65,4.25])
step(3, 'Deploy dan buka domain utama project. Jika menggunakan domain sendiri, hubungkan domain dan pastikan HTTPS berfungsi.')
h('Autentikasi Supabase')
step(4, 'Buka Authentication > Sign In / Providers. Matikan Allow new users to sign up dan anonymous sign-in. Aktifkan provider Email; matikan provider lain yang tidak digunakan.')
step(5, 'Pada provider Email, atur minimum password 12 karakter. Tinjau masa berlaku link email dan batas permintaan Auth; contoh masa berlaku link adalah satu jam.')
step(6, 'Buka Authentication > URL Configuration. Isi Site URL dan tambahkan Redirect URL yang sama dengan alamat admin berikut, memakai domain asli pembeli:')
code('https://domain-pembeli.vercel.app/admin/')
p('Gunakan domain utama yang stabil, bukan alamat preview yang berubah setiap deployment. Sertakan /admin/ dan garis miring terakhir. Hindari wildcard redirect untuk produksi. Simpan pengaturan.')

page('4 Atur email undangan dan pemulihan')
p('SMTP adalah layanan pengirim email. Email pengirim boleh milikmu; email penerima dan login tetap milik pembeli. Atur SMTP pada setiap project Supabase.')
h('Pilihan Gmail untuk tahap awal')
step(1, 'Login ke christozuang@gmail.com. Buka myaccount.google.com/security dan aktifkan Verifikasi 2 Langkah.')
step(2, 'Buka myaccount.google.com/apppasswords. Buat sandi aplikasi, misalnya Supabase Portfolio Budi. Salin sandi 16 karakter yang diberikan Google.')
step(3, 'Buka Authentication > Email > SMTP Settings di Supabase. Aktifkan Custom SMTP dan isi:')
table(['Kolom', 'Nilai'], [('Sender email','christozuang@gmail.com'),('Sender name','Portfolio Admin'),('Host','smtp.gmail.com'),('Port','587'),('Username','christozuang@gmail.com'),('Password','App Password Google tanpa spasi dan tanda petik')], [1.75,5.15])
p('Gunakan App Password, bukan password utama Gmail. Jika menu App Passwords tidak tersedia, periksa status Verifikasi 2 Langkah dan pembatasan keamanan akun. Jangan membagikan sandi ini kepada pembeli atau memasukkannya ke repository.')
p('Satu App Password dapat dipakai pada beberapa project. Sandi terpisah per project memudahkan pencabutan akses. Semua project tetap berbagi batas pengiriman akun Gmail yang sama; membuat sandi baru tidak menambah kuota. Batas berlaku menurut kebijakan Google, sehingga Gmail sebaiknya digunakan untuk pengujian atau volume kecil.')
h('Template email dan layanan untuk penjualan rutin')
step(4, 'Pada Email Templates, periksa Invite user dan Reset password. Pertahankan tautan bawaan seperti berikut, lalu simpan:')
code('<a href="{{ .ConfirmationURL }}">Continue</a>')
p('Jangan mengganti ConfirmationURL dengan alamat /admin/ langsung atau template PKCE/token hash. Implementasi saat ini menerima callback bawaan Supabase untuk invite dan recovery.')
p('Untuk pengiriman rutin, gunakan penyedia email transaksional seperti Brevo atau Resend dengan domain terverifikasi. Ikuti host, port, kredensial, serta kuota dari penyedia tersebut. Email bawaan Supabase memiliki pembatasan penerima dan pengiriman.')

page('5 Undang pembeli dan aktifkan akun')
h('Langkah penjual')
step(1, 'Buka Authentication > Users > Invite user. Masukkan email pembeli yang benar, lalu kirim undangan.')
step(2, 'Salin UUID akun yang baru diundang dari daftar Users. UUID adalah ID akun, bukan alamat email.')
step(3, 'Buka SQL Editor project pembeli, lalu jalankan:')
code('''insert into public.admin_users (user_id)
values ('GANTI_DENGAN_UUID_PEMBELI')
on conflict (user_id) do nothing;''')
p('Tanda petik tunggal tetap digunakan. Ganti hanya GANTI_DENGAN_UUID_PEMBELI dengan UUID asli. Pastikan UUID berasal dari project Supabase yang sedang dikerjakan.')
step(4, 'Periksa Table Editor > admin_users. Pastikan UUID pembeli sudah ada, baru minta pembeli membuka email undangannya.')
p('Undangan dikirim sebelum pemberian izin selesai. Jika pembeli membuka link terlalu cepat, aplikasi akan menolak akses sampai UUID diberi izin. Jangan mengaktifkan signup publik sebagai solusi.')
h('Langkah pembeli')
step(1, 'Buka email undangan dan klik tombol aktivasi. Jika tidak terlihat, periksa folder spam.')
step(2, 'Pada halaman Activate your portfolio, buat password minimal 12 karakter dan isi konfirmasinya.')
step(3, 'Klik Save password, lalu login dengan email dan password tersebut.')
step(4, 'Untuk login selanjutnya, buka alamat berikut dengan domain portfolio sendiri:')
code('https://domain-pembeli.vercel.app/admin/')
p('Link undangan tidak digunakan setiap kali login. Jika link kedaluwarsa atau sudah dipakai, gunakan Forgot password? untuk meminta link baru bagi akun yang sudah dibuat.')
h('Batas akses')
p('Memiliki akun Auth belum berarti menjadi admin. Izin pengelolaan berasal dari admin_users dan diperiksa oleh RLS pada database serta Storage. Pembeli tidak menerima akses dashboard Supabase, Vercel, atau kredensial pengirim email melalui alur ini.')

page('6 Uji dan lakukan serah terima')
h('Checklist sebelum website diserahkan')
for text in [
'Website publik dan /admin/ terbuka melalui HTTPS pada domain yang benar.',
'Email undangan diterima; pembeli berhasil membuat password dan login.',
'Ubah satu konten percobaan, simpan, muat ulang, dan periksa hasilnya di website publik.',
'Upload gambar percobaan berhasil dan dapat ditampilkan di website.',
'Forgot password? mengirim email; password dapat diganti dan dipakai untuk login.',
'Signup publik tetap mati dan akun tanpa admin_users ditolak dari editor.',
'Tampilan di HP diperiksa; semua teks, gambar, kontak, dan tautan contoh sudah diganti.',
'Jika memakai akun pengujian, hapus izinnya dari admin_users setelah selesai.'
]: p('[ ] ' + text)
h('Pesan yang dapat dikirim kepada pembeli')
p('Portfolio kamu sudah siap. Berikut alamat website dan halaman pengelolaannya.')
code('''Website: https://domain-pembeli.vercel.app/
Admin: https://domain-pembeli.vercel.app/admin/
Email login: email-pembeli@example.com''')
p('Silakan buka email undangan untuk membuat password sendiri. Setelah aktivasi, login melalui halaman admin. Jika lupa password, pilih Forgot password?. Jangan membagikan link aktivasi atau password kepada orang lain.')
h('Catatan internal setiap penjualan')
table(['Informasi', 'Isi untuk pembeli ini'], [
('Nama dan email pembeli','________________________________'),
('Repository dan project Vercel','________________________________'),
('Project Supabase dan UUID admin','________________________________'),
('Domain dan tanggal serah terima','________________________________'),
('Paket layanan dan jadwal perpanjangan','________________________________')], [2.85,4.05])
p('Simpan rahasia layanan di pengelola password, bukan dalam catatan serah terima. Tentukan jadwal backup konten dan aset serta prosedur pemulihan sesuai layanan yang disepakati.')

page('7 Atasi masalah yang umum terjadi')
h('Vercel mencari folder public')
p('Jika muncul No Output Directory named "public", buka Settings > Build and Deployment. Atur Build Command ke npm run build dan Output Directory ke dist. Simpan lalu redeploy. Log build harus menunjukkan hasil di folder dist.')
h('Undangan gagal dikirim')
p('Jika muncul Error sending invite email, buka Logs dan filter Log Type ke Auth. Cari error pada waktu undangan dikirim, lalu baca Raw JSON. Log API Gateway dengan status 500 saja belum menjelaskan penyebab SMTP.')
table(['Detail error', 'Pemeriksaan'], [
('535 atau password not accepted','Cocokkan username, sender email, dan App Password Gmail.'),
('534 atau application password required','Gunakan App Password Google, bukan password utama Gmail.'),
('Timeout atau connection refused','Periksa host, port, dan ketersediaan layanan SMTP.'),
('Rate limit atau quota exceeded','Tunggu batas pulih; periksa kuota Supabase dan penyedia email.')], [2.5,4.4])
h('Akun belum mendapat izin atau link tidak berlaku')
p('Untuk pesan Administrator access has not been granted, cocokkan UUID Auth dengan admin_users pada project yang benar. Untuk link invalid, expired, atau sudah dipakai, minta link pemulihan melalui Forgot password?. Jangan mengubah role melalui metadata pengguna.')
h('Admin hanya menampilkan preview atau gagal menyimpan')
p('Preview menunjukkan Supabase belum dikonfigurasi. Periksa config.js yang ikut dalam deployment dan lakukan deploy ulang setelah perubahan. Jika simpan gagal, periksa record site, isi JSON lengkap, membership admin_users, dan RLS. Jika muncul save conflict, muat ulang konten sebelum mencoba lagi.')
h('Mencabut akses pembeli')
p('Untuk menghentikan akses editor, hapus hanya baris UUID pembeli yang dimaksud dari admin_users. Kelola penonaktifan akun di Authentication jika diperlukan. Verifikasi project dan UUID sebelum tindakan tersebut. Jangan menghapus database, Storage, atau deployment tanpa kesepakatan dengan pembeli.')
p('Saat mengganti domain, perbarui Site URL dan Redirect URLs, lalu uji ulang undangan serta pemulihan password. Jika kredensial SMTP dicabut, perbarui semua project yang menggunakannya.')

footer = s.footer.paragraphs[0]
footer.alignment = 2
run = footer.add_run('Panduan penjualan portfolio  |  ')
run.font.size = Pt(9)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
